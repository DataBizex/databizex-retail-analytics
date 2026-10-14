import copy
import hashlib
import json
import shutil

import numpy as np
import pandas as pd
import pytest

from app.analytics.kpis import prepare_sales_frame, validity_summary
from app.data import fixes
from app.data.change_log import SOURCE_ROWS_CAP, ChangeLog
from app.data.fixes import FixAction, FixError
from app.data.loader import SOURCE_ROW, load_file, with_source_rows
from app.data.rules import Rule, default_rules, default_rules_path, expand_rules, quality_score, validate
from app.paths import sample_data_path

RULES = default_rules()


def _frame():
    """Index = source line numbers, like the working copy.

    2 clean | 3 invalid date, "germany " variant, quantity -1 | 4 country missing, quantity "two",
    discount 150 | 5 exact duplicate of 2 | 6 "GERMANY", quantity missing | 7 price 5 below cost 8
    """
    rows = [
        ["A1", "2024-01-01", "Germany", "Chair", "Furniture", "2", "100", "10", "40"],
        ["A2", "2024-13-45", "germany ", "Lamp", "Lighting", "-1", "50", "0", "20"],
        ["A3", "2024-02-01", None, "Lamp", "Lighting", "two", "50", "150", "20"],
        ["A1", "2024-01-01", "Germany", "Chair", "Furniture", "2", "100", "10", "40"],
        ["A4", "2024-02-03", "GERMANY", "Mug", "Kitchen", None, "10", "0", "4"],
        ["A5", "2024-02-04", "France", "Mug", "Kitchen", "3", "5", "0", "8"],
    ]
    columns = ["order_id", "order_date", "country", "product", "category", "quantity", "unit_price",
               "discount_pct", "unit_cost"]
    return pd.DataFrame(rows, columns=columns, index=pd.Index([2, 3, 4, 5, 6, 7], name=SOURCE_ROW))


# The 6-row fixture is too small for expand_rules to treat country as categorical, so the
# per-column text variants rule is built the way expand_rules builds it.
TEXT_VARIANTS_COUNTRY = Rule(name="Text variants: country", column="country", type="text_variants",
                             severity="warning")


def _rule(name, df=None):
    if name == TEXT_VARIANTS_COUNTRY.name:
        return TEXT_VARIANTS_COUNTRY
    df = _frame() if df is None else df
    return {r.name: r for r in expand_rules(RULES, df)}[name]


@pytest.fixture
def df():
    return _frame()


def _apply_twice(df, rule, action, params=None):
    """(result, entry) of one fix; also checks the input is untouched and a second run changes nothing."""
    original = copy.deepcopy(df)
    out, entry = fixes.run_fix(df, rule, action, params, RULES)
    pd.testing.assert_frame_equal(df, original)  # the input (original / working copy) is never modified
    again, second = fixes.run_fix(out, rule, action, params, RULES)
    pd.testing.assert_frame_equal(again, out)  # idempotent
    assert second.rows_affected == 0 and second.cells_changed == 0 and second.source_rows == []
    return out, entry


# ---------------------------------------------------------------------- one group per action
def test_exclude_rows(df):
    rule = _rule("Negative values: quantity")
    out, entry = _apply_twice(df, rule, FixAction.EXCLUDE_ROWS)
    assert list(out.index) == [2, 4, 5, 6, 7]
    assert (entry.rule, entry.action, entry.columns) == ("Negative values: quantity", "EXCLUDE_ROWS", ["quantity"])
    assert (entry.rows_affected, entry.cells_changed, entry.source_rows) == (1, 0, [3])
    assert entry.samples == [{"source_row": 3, "column": "quantity", "before": "quantity=-1",
                              "after": "(row excluded)"}]


def test_set_missing_keeps_the_rows(df):
    rule = _rule("Negative values: quantity")
    out, entry = _apply_twice(df, rule, FixAction.SET_MISSING)
    assert list(out.index) == list(df.index)  # no row removed
    assert out.at[3, "quantity"] is None and out.drop(columns="quantity").equals(df.drop(columns="quantity"))
    assert (entry.rows_affected, entry.cells_changed, entry.source_rows) == (1, 1, [3])
    assert entry.samples == [{"source_row": 3, "column": "quantity", "before": "-1", "after": "(missing)"}]


def test_set_missing_on_a_float_and_a_date_column():
    df = pd.DataFrame({"quantity": [1.0, -2.0, 3.0], "order_date": pd.to_datetime(["2024-01-01"] * 3)},
                      index=pd.Index([2, 3, 4], name=SOURCE_ROW))
    rule = Rule(name="neg", column="quantity", type="range", params={"min": 0})
    out, entry = fixes.set_missing(df, rule)
    assert np.isnan(out.at[3, "quantity"]) and out["quantity"].dtype == float and entry.cells_changed == 1


def test_map_values(df):
    rule = _rule("Text variants: country")
    candidates = fixes.mapping_candidates(df, rule)
    # Only the failing variants are listed, each with the normalized form as a suggestion.
    assert sorted(candidates) == [("GERMANY", 1, "Germany"), ("germany ", 1, "Germany")]
    mapping = {"germany ": "Germany", "GERMANY": "Germany", "Germany": ""}  # empty target: ignored
    out, entry = _apply_twice(df, rule, FixAction.MAP_VALUES, {"mapping": mapping})
    assert list(out["country"]) == ["Germany", "Germany", None, "Germany", "Germany", "France"]
    assert (entry.rows_affected, entry.cells_changed, entry.source_rows) == (2, 2, [3, 6])
    assert entry.parameters == {"mapping": {"germany ": "Germany", "GERMANY": "Germany"}}


def test_map_values_matches_case_and_spacing_and_leaves_unmapped_values():
    df = pd.DataFrame({"city": ["NY", " ny", "New York", "N.Y.", "new  york", "Boston", None]},
                      index=pd.Index(range(2, 9), name=SOURCE_ROW))
    rule = Rule(name="city values", column="city", type="allowed_values", params={"values": ["New York"]})
    out, entry = fixes.map_values(df, rule, {"ny": "New York", "N.Y.": "New York", "new york": "New York",
                                             "Boston": "  "})
    assert list(out["city"]) == ["New York", "New York", "New York", "New York", "New York", "Boston", None]
    assert entry.cells_changed == 4  # "New York" was already right; Boston has an empty target


def test_fill_constant(df):
    rule = _rule("Missing values: quantity")
    out, entry = _apply_twice(df, rule, FixAction.FILL_CONSTANT, {"value": "1"})
    assert out.at[6, "quantity"] == "1"  # a text column of numbers keeps text
    assert (entry.rows_affected, entry.cells_changed, entry.source_rows) == (1, 1, [6])
    assert entry.parameters == {"value": "1", "kind": "numeric"}
    date_rule = _rule("Invalid dates: order_date")
    out, entry = fixes.fill_constant(df, date_rule, "2024-03-15")
    assert out.at[3, "order_date"] == "2024-03-15" and entry.cells_changed == 1


@pytest.mark.parametrize("rule_name, value, message", [
    ("Missing values: quantity", "lots", "is not a number"),
    ("Invalid dates: order_date", "soon", "is not a valid date"),
    ("Negative values: quantity", "-3", "would still fail the rule"),
    ("Missing values: quantity", "  ", "Enter a value"),
])
def test_fill_constant_rejects_a_wrong_value_and_changes_nothing(df, rule_name, value, message):
    original = copy.deepcopy(df)
    with pytest.raises(FixError, match=message):
        fixes.run_fix(df, _rule(rule_name), FixAction.FILL_CONSTANT, {"value": value}, RULES)
    pd.testing.assert_frame_equal(df, original)


def test_fill_constant_text_column(df):
    out, entry = fixes.fill_constant(df, _rule("Missing values: country"), "  Unknown ")
    assert out.at[4, "country"] == "Unknown" and entry.parameters["kind"] == "text"


def test_trim_and_case(df):
    rule = _rule("Text variants: country")
    out, entry = _apply_twice(df, rule, FixAction.TRIM_AND_CASE, {"case": "title"})
    assert list(out["country"]) == ["Germany", "Germany", None, "Germany", "Germany", "France"]
    assert (entry.rows_affected, entry.cells_changed, entry.source_rows) == (2, 2, [3, 6])
    out, entry = fixes.trim_and_case(df, rule, "keep")
    assert out.at[3, "country"] == "germany" and out.at[6, "country"] == "GERMANY" and entry.cells_changed == 1
    with pytest.raises(FixError):
        fixes.trim_and_case(df, rule, "shouting")


def test_fix_duplicates(df):
    rule = _rule("Exact duplicate rows")
    out, entry = _apply_twice(df, rule, FixAction.FIX_DUPLICATES)
    assert list(out.index) == [2, 3, 4, 6, 7]
    assert (entry.rows_affected, entry.source_rows) == (1, [5])
    assert entry.parameters["kept_row_for_removed"] == {5: 2}
    assert entry.samples[0]["after"] == "(removed; row 2 kept)"
    keyed = _rule("Duplicate order_id")
    out, entry = fixes.fix_duplicates(df, keyed)
    assert entry.columns == ["order_id"] and entry.source_rows == [5]


def test_source_file_is_never_modified(tmp_path):
    copy_path = tmp_path / "demo.csv"
    shutil.copyfile(sample_data_path(), copy_path)
    digest = hashlib.sha256(copy_path.read_bytes()).hexdigest()
    loaded = load_file(copy_path)
    working = with_source_rows(loaded.df, loaded.source_rows)
    original = copy.deepcopy(loaded.df)
    rules = {r.name: r for r in expand_rules(RULES, working)}
    for name, action, params in [("Negative values: quantity", FixAction.SET_MISSING, None),
                                 ("Missing values: country", FixAction.FILL_CONSTANT, {"value": "Unknown"}),
                                 ("Exact duplicate rows", FixAction.FIX_DUPLICATES, None),
                                 ("Invalid numeric: quantity", FixAction.EXCLUDE_ROWS, None)]:
        working, _entry = fixes.run_fix(working, rules[name], action, params, RULES)
    assert hashlib.sha256(copy_path.read_bytes()).hexdigest() == digest
    pd.testing.assert_frame_equal(loaded.df, original)


# ---------------------------------------------------------------------- suggestions
def _default_rule_items():
    return json.loads(default_rules_path().read_text(encoding="utf-8"))["rules"]


@pytest.mark.parametrize("item", _default_rule_items(), ids=lambda item: item["name"])
def test_every_default_rule_has_suggested_actions(item):
    rule = next(r for r in RULES if r.name == item["name"])
    assert rule.type in fixes.ACTIONS_BY_TYPE, f"no fix mapping for rule type {rule.type}"
    actions = fixes.suggested_actions(rule, object)
    if rule.type == "outlier":
        assert actions == []  # outliers are warnings only: no automatic removal
    elif rule.column == "*":
        assert set(actions) <= fixes.ROW_ACTIONS
    else:
        assert actions and all(isinstance(a, FixAction) for a in actions)
    for action in actions:
        assert fixes.DESCRIPTIONS[action]


def test_default_rule_set_has_25_rules():
    assert len(_default_rule_items()) == 25


def test_suggestions_depend_on_column_type_and_unknown_types_fall_back():
    variants = Rule(name="v", column="country", type="text_variants")
    assert FixAction.MAP_VALUES in fixes.suggested_actions(variants, object)
    assert fixes.suggested_actions(Rule(name="d", column="q", type="dtype", params={"kind": "numeric"}),
                                   np.dtype("float64")) == [FixAction.SET_MISSING, FixAction.FILL_CONSTANT,
                                                            FixAction.EXCLUDE_ROWS]
    assert fixes.suggested_actions(Rule(name="x", column="q", type="new_type")) == fixes.FALLBACK_ACTIONS


def test_actions_not_offered_are_refused(df):
    with pytest.raises(FixError, match="not offered"):
        fixes.run_fix(df, _rule("Outliers (IQR): quantity"), FixAction.EXCLUDE_ROWS, None, RULES)


def test_what_it_will_do_uses_real_counts(df):
    _out, entry = fixes.run_fix(df, _rule("Negative values: quantity"), FixAction.SET_MISSING, None, RULES)
    assert fixes.what_it_will_do(entry) == "1 value(s) in quantity will be set to missing; the rows stay."


# ---------------------------------------------------------------------- scores
def test_scores_match_a_fresh_recomputation(df):
    out, entry = fixes.run_fix(df, _rule("Exact duplicate rows"), FixAction.FIX_DUPLICATES, None, RULES)
    for frame, score, valid in ((df, entry.quality_score_before, entry.valid_rows_before),
                                (out, entry.quality_score_after, entry.valid_rows_after)):
        summary, flags = validate(frame, RULES)
        assert score == pytest.approx(quality_score(RULES, summary, flags).overall)
        assert valid == validity_summary(prepare_sales_frame(frame)).kpi_rows
    assert entry.valid_rows_after == entry.valid_rows_before - 1  # the duplicate was a valid row


# ---------------------------------------------------------------------- change log and undo
def test_undo_restores_exactly_after_three_fixes(df):
    log = ChangeLog()
    states = [df]
    steps = [("Negative values: quantity", FixAction.SET_MISSING, None),
             ("Text variants: country", FixAction.TRIM_AND_CASE, {"case": "title"}),
             ("Exact duplicate rows", FixAction.FIX_DUPLICATES, None)]
    working = df
    for name, action, params in steps:
        new, entry = fixes.run_fix(working, _rule(name, working), action, params, RULES)
        log.add(entry, working)
        working = new
        states.append(working)
    assert [e.id for e in log.entries] == [1, 2, 3]
    for expected in reversed(states[:-1]):
        entry, restored = log.undo_last()
        pd.testing.assert_frame_equal(restored, expected)
    assert len(log) == 0 and log.undo_last() is None


def test_snapshot_stack_is_bounded(df):
    log = ChangeLog(snapshot_limit=10)
    rule = _rule("Negative values: quantity")
    for i in range(12):
        _out, entry = fixes.run_fix(df, rule, FixAction.SET_MISSING, None, RULES)
        log.add(entry, df.assign(step=i))
    assert len(log) == 12 and log.undoable_count() == 10
    for i in reversed(range(2, 12)):
        _entry, state = log.undo_last()
        assert state["step"].iloc[0] == i
    assert not log.can_undo() and log.undo_last() is None and len(log) == 2  # older fixes stay logged
    log.clear()
    assert len(log) == 0 and log.undoable_count() == 0


def test_entry_metadata(df):
    log = ChangeLog()
    _out, entry = fixes.run_fix(df, _rule("Negative values: quantity"), FixAction.SET_MISSING, None, RULES)
    log.add(entry, df)
    assert entry.id == 1 and "T" in entry.timestamp and entry.kind == "fix"


def test_truncation_flag_above_5000_rows():
    n = SOURCE_ROWS_CAP + 1234
    df = pd.DataFrame({"quantity": [-1.0] * n}, index=pd.Index(range(2, n + 2), name=SOURCE_ROW))
    rule = Rule(name="neg", column="quantity", type="range", params={"min": 0})
    _out, entry = fixes.set_missing(df, rule)
    assert entry.truncated and entry.source_rows_total == n == entry.rows_affected
    assert len(entry.source_rows) == SOURCE_ROWS_CAP and entry.source_rows[0] == 2
    assert len(entry.samples) == 20
    _out, small = fixes.set_missing(df.head(10), rule)
    assert not small.truncated


def test_change_log_exports(df):
    log = ChangeLog()
    log.preparation = ["Import: semicolon delimiter"]
    out, entry = fixes.run_fix(df, _rule("Negative values: quantity"), FixAction.SET_MISSING, None, RULES)
    log.add(entry, df)
    frame = log.to_dataframe()
    assert list(frame["action"]) == ["SET_MISSING"] and frame.loc[0, "source_rows"] == "3"
    csv = log.to_csv_text()
    assert csv.splitlines()[0].startswith("id,time,kind,rule,action") and "Negative values: quantity" in csv
    md = log.to_markdown("demo.csv")
    assert md.startswith("# DataBizex change log") and "- Source file: demo.csv" in md
    assert "## 1. Negative values: quantity: SET_MISSING" in md and "| 3 | quantity | -1 | (missing) |" in md
    assert "Preparation (before any fix): Import: semicolon delimiter" in md
    assert "No fixes applied." in ChangeLog().to_markdown("x.csv")


# ---------------------------------------------------------------------- exports
def _log_with_one_fix(df):
    log = ChangeLog()
    _out, entry = fixes.run_fix(df, _rule("Negative values: quantity"), FixAction.SET_MISSING, None, RULES)
    log.add(entry, df)
    return log


def test_export_change_log_csv_and_markdown(df, tmp_path):
    from app.exports.exporter import CSV_ENCODING, export_change_log

    log = _log_with_one_fix(df)
    csv_path = export_change_log(log, tmp_path / "log.csv", "demo.csv")
    raw = csv_path.read_bytes()
    assert raw.startswith("﻿".encode("utf-8")) and CSV_ENCODING == "utf-8-sig"  # like the other CSVs
    assert pd.read_csv(csv_path, encoding=CSV_ENCODING).loc[0, "rule"] == "Negative values: quantity"
    md_path = export_change_log(log, tmp_path / "log.md", "demo.csv")
    md = md_path.read_bytes()
    assert not md.startswith(b"\xef\xbb\xbf")  # Markdown: UTF-8 without BOM
    assert md.decode("utf-8").startswith("# DataBizex change log")


def test_excel_change_log_sheet_only_when_not_empty(df, tmp_path):
    from openpyxl import load_workbook

    from app.exports.exporter import export_excel

    with_log = export_excel(df, tmp_path / "with.xlsx", "working_data", _log_with_one_fix(df))
    without = export_excel(df, tmp_path / "without.xlsx", "working_data", ChangeLog())
    assert load_workbook(with_log).sheetnames == ["working_data", "Change log"]
    assert load_workbook(without).sheetnames == ["working_data"]
    assert pd.read_excel(with_log, sheet_name="Change log").loc[0, "action"] == "SET_MISSING"


def test_powerbi_package_change_log_only_when_not_empty(kpi_df, tmp_path):
    from app.exports.exporter import generate_powerbi_package

    working = with_source_rows(kpi_df, list(range(2, 2 + len(kpi_df))))
    rule = {r.name: r for r in expand_rules(RULES, working)}["Negative values: quantity"]
    log = ChangeLog()
    _out, entry = fixes.run_fix(working, rule, FixAction.SET_MISSING, None, RULES)
    log.add(entry, working)
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    names = {p.name for p in generate_powerbi_package(working, tmp_path / "a", "kpi.csv", change_log=log)}
    assert {"change_log.csv", "change_log.md"} <= names
    assert "Fixes applied to the working copy: 1" in (tmp_path / "a" / "package_info.txt").read_text("utf-8")
    names = {p.name for p in generate_powerbi_package(working, tmp_path / "b", "kpi.csv", change_log=ChangeLog())}
    assert not names & {"change_log.csv", "change_log.md"}
    assert "Fixes applied" not in (tmp_path / "b" / "package_info.txt").read_text("utf-8")
